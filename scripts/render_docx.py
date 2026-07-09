#!/usr/bin/env python3
from __future__ import annotations

import argparse
import sys
from pathlib import Path

import yaml
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from build_metadata import build_date, copy_named_artifact, named_artifact_path


BODY_FONT = "Arial"
BODY_SIZE = Pt(10.5)
CONTACT_SIZE = Pt(9.5)
TITLE_SIZE = Pt(18)
SECTION_SIZE = Pt(11)
ROLE_SIZE = Pt(10.5)
SECTION_COLOR = RGBColor(31, 58, 95)
BORDER_COLOR = "8EAADB"


def load_resume(path: Path) -> dict:
    with path.open("r", encoding="utf-8") as handle:
        data = yaml.safe_load(handle)
    if not isinstance(data, dict):
        raise ValueError(f"{path}: expected a YAML mapping at the document root")
    return data


def set_paragraph_spacing(paragraph, *, before: int = 0, after: int = 0, line: float = 1.0) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def set_run_font(run, *, bold: bool | None = None, size=None) -> None:
    run.font.name = BODY_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold


def configure_styles(document: Document) -> None:
    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = BODY_SIZE
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.05

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = styles[style_name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.font.bold = True

    styles["Heading 1"].font.size = Pt(12)
    styles["Heading 1"].font.color.rgb = SECTION_COLOR
    styles["Heading 1"].paragraph_format.space_before = Pt(11)
    styles["Heading 1"].paragraph_format.space_after = Pt(3)

    styles["Heading 2"].font.size = SECTION_SIZE
    styles["Heading 2"].font.color.rgb = None
    styles["Heading 2"].paragraph_format.space_before = Pt(7)
    styles["Heading 2"].paragraph_format.space_after = Pt(2)

    styles["Heading 3"].font.size = ROLE_SIZE
    styles["Heading 3"].font.color.rgb = None
    styles["Heading 3"].paragraph_format.space_before = Pt(4)
    styles["Heading 3"].paragraph_format.space_after = Pt(2)

    bullet = styles["List Bullet"]
    bullet.font.name = BODY_FONT
    bullet._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    bullet.font.size = BODY_SIZE
    bullet.paragraph_format.left_indent = Inches(0.25)
    bullet.paragraph_format.first_line_indent = Inches(-0.15)
    bullet.paragraph_format.space_after = Pt(3)
    bullet.paragraph_format.line_spacing = 1.05

    if "Resume Contact" not in styles:
        contact = styles.add_style("Resume Contact", WD_STYLE_TYPE.PARAGRAPH)
        contact.base_style = normal
    contact = styles["Resume Contact"]
    contact.font.name = BODY_FONT
    contact._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    contact.font.size = CONTACT_SIZE
    contact.paragraph_format.space_after = Pt(6)
    contact.paragraph_format.line_spacing = 1.0


def configure_page(document: Document) -> None:
    section = document.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.right_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def add_bottom_border(paragraph, *, color: str = BORDER_COLOR, size: str = "6", space: str = "1") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)

    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)

    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)


def add_paragraph(document: Document, text: str = "", style: str | None = None, *, bold: bool | None = None):
    paragraph = document.add_paragraph(style=style)
    if text:
        run = paragraph.add_run(text)
        set_run_font(run, bold=bold)
    return paragraph


def add_title_block(document: Document, profile: dict, headline: str) -> None:
    name = add_paragraph(document)
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(name, after=1, line=1.0)
    run = name.add_run(profile["name"])
    set_run_font(run, bold=True, size=TITLE_SIZE)

    contact_parts = [
        profile["location"],
        profile["email"],
        profile["phone"],
        profile["linkedin"],
        *profile.get("websites", []),
    ]
    contact = add_paragraph(document, " | ".join(contact_parts), "Resume Contact")
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER

    headline_paragraph = add_paragraph(document)
    headline_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(headline_paragraph, after=8, line=1.0)
    run = headline_paragraph.add_run(headline)
    set_run_font(run, bold=True, size=BODY_SIZE)


def add_section_heading(document: Document, text: str) -> None:
    paragraph = add_paragraph(document, text, "Heading 1")
    keep_with_next(paragraph)
    add_bottom_border(paragraph)


def add_role_heading(document: Document, title: str, start: str, end: str) -> None:
    paragraph = add_paragraph(document, f"{title} | {start} - {end}", "Heading 3")
    keep_with_next(paragraph)


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.first_line_indent = Inches(-0.15)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.05
    bullet = paragraph.add_run("•\t")
    set_run_font(bullet)
    run = paragraph.add_run(text)
    set_run_font(run)


def add_summary(document: Document, summary: list[str]) -> None:
    add_section_heading(document, "Summary")
    for item in summary:
        add_bullet(document, item)


def add_skills(document: Document, skills: list[str]) -> None:
    add_section_heading(document, "Skills")
    for skill in skills:
        add_bullet(document, skill)


def add_projects(document: Document, projects: list[dict]) -> None:
    if not projects:
        return
    add_section_heading(document, "Projects")
    for project in projects:
        heading = document.add_paragraph(style="Heading 3")
        name = heading.add_run(project["name"])
        set_run_font(name, bold=True)
        if project.get("url"):
            url = heading.add_run(f" ({project['url']})")
            set_run_font(url, bold=False)
        keep_with_next(heading)
        description = add_paragraph(document, project["description"])
        set_paragraph_spacing(description, after=5, line=1.05)


def add_experience(document: Document, experience: list[dict]) -> None:
    add_section_heading(document, "Professional Experience")
    for job in experience:
        company_text = job["company"]
        if job.get("location"):
            company_text = f"{company_text} - {job['location']}"
        company = add_paragraph(document, company_text, "Heading 2")
        keep_with_next(company)
        for role in job["roles"]:
            add_role_heading(document, role["title"], role["start"], role["end"])
            for bullet in role["bullets"]:
                add_bullet(document, bullet)


def education_line(item: dict) -> str:
    degree = item["degree"]
    if item.get("field"):
        degree = f"{degree} in {item['field']}"
    if item.get("start") or item.get("end"):
        return f"{degree} | {item.get('start', '')} - {item.get('end', '')}".strip()
    return degree


def add_education(document: Document, education: list[dict]) -> None:
    add_section_heading(document, "Education")
    for item in education:
        school = add_paragraph(document, f"{item['institution']} - {item['location']}", "Heading 3")
        keep_with_next(school)
        set_paragraph_spacing(school, before=4, after=0, line=1.0)
        line = add_paragraph(document, education_line(item))
        line.paragraph_format.left_indent = Inches(0.15)
        set_paragraph_spacing(line, after=5, line=1.0)


def add_document_language(document: Document) -> None:
    settings = document.settings._element
    lang = settings.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        settings.append(lang)
    lang.set(qn("w:val"), "en-US")


def set_core_properties(document: Document, date_string: str) -> None:
    properties = document.core_properties
    properties.version = date_string
    properties.keywords = f"resume-build-date:{date_string}"


def render_docx(resume_path: Path, output_path: Path, date_override: str | None) -> int:
    resume = load_resume(resume_path)
    date_string = build_date(date_override)
    document = Document()
    configure_page(document)
    configure_styles(document)
    add_document_language(document)
    set_core_properties(document, date_string)

    add_title_block(document, resume["profile"], resume["headline"])
    add_summary(document, resume["summary"])
    add_skills(document, resume["skills"])
    add_projects(document, resume.get("projects", []))
    add_experience(document, resume["experience"])
    add_education(document, resume["education"])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    named_docx = named_artifact_path(output_path.parent, resume, date_string, ".docx")
    copy_named_artifact(output_path, named_docx)
    print(f"Rendered {output_path}")
    print(f"Rendered {named_docx}")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description="Render ATS-safe DOCX from resume YAML.")
    parser.add_argument("--resume", type=Path, default=Path("resume.yaml"))
    parser.add_argument("--output", type=Path, default=Path("output/ats.docx"))
    parser.add_argument("--date", default=None, help="UTC build date override in YYYYMMDD form")
    args = parser.parse_args()

    try:
        return render_docx(args.resume, args.output, args.date)
    except (OSError, ValueError, KeyError, yaml.YAMLError) as exc:
        print(f"DOCX render error: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
