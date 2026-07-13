#!/usr/bin/env python3
from __future__ import annotations

import argparse
import sys
import zipfile
from pathlib import Path

import yaml
from docx import Document
from docx.oxml.ns import qn

from build_metadata import default_artifact_path
from check_parse import (
    as_string_list,
    check_forbidden,
    check_order,
    check_required_exact,
    check_role_boundaries,
    check_education_degree_fields,
    expected_document_order,
    check_role_noun_starts,
    check_yaml_content_present,
    load_expectations,
    load_resume,
    role_heading,
)


STANDARD_DOCX_HEADINGS = {"Summary", "Skills", "Projects", "Professional Experience", "Education"}


def docx_text(document: Document) -> str:
    return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text)


def read_document_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        with archive.open("word/document.xml") as handle:
            return handle.read().decode("utf-8")


def check_docx_role_headings(document: Document, resume: dict) -> list[str]:
    errors: list[str] = []
    paragraphs = [(paragraph.text.strip(), paragraph.style.name) for paragraph in document.paragraphs]
    cursor = 0

    for job_index, job in enumerate(resume.get("experience", [])):
        if not isinstance(job, dict):
            continue
        legacy_heading = (f"{job.get('company')} - {job.get('location')}", "Heading 2")
        if legacy_heading in paragraphs:
            errors.append(
                f"experience[{job_index}] uses a shared employer heading; every role must be a flat entry"
            )
        for role_index, role in enumerate(job.get("roles", [])):
            if not isinstance(role, dict):
                continue
            try:
                expected = (role_heading(str(job.get("company")), str(job.get("location")), role), "Heading 3")
            except KeyError as exc:
                errors.append(
                    f"experience[{job_index}].roles[{role_index}] is missing key needed for DOCX heading: {exc}"
                )
                continue
            try:
                found = paragraphs.index(expected, cursor)
            except ValueError:
                errors.append(
                    f"complete DOCX role heading missing for experience[{job_index}].roles[{role_index}]: "
                    f"{expected[0]!r}"
                )
                continue
            cursor = found + 1

    return errors


def check_docx_structure(path: Path, document: Document) -> list[str]:
    errors: list[str] = []
    if document.tables:
        errors.append(f"DOCX contains {len(document.tables)} table(s); ATS output should be single-column prose")
    if document.inline_shapes:
        errors.append(f"DOCX contains {len(document.inline_shapes)} inline image/shape(s)")

    xml = read_document_xml(path)
    forbidden_xml_markers = {
        "drawing": "<w:drawing",
        "legacy picture": "<w:pict",
        "text box": "txbxContent",
        "table": "<w:tbl",
    }
    for label, marker in forbidden_xml_markers.items():
        if marker in xml:
            errors.append(f"DOCX XML contains {label} markup: {marker}")

    if "<w:numPr>" in xml:
        errors.append("DOCX contains Word numbering properties; ATS DOCX bullets must be literal glyph paragraphs")

    literal_bullets = 0
    bordered_headings = set()
    for paragraph in document.paragraphs:
        stripped = paragraph.text.strip()
        p_pr = paragraph._p.pPr
        has_paragraph_border = p_pr is not None and p_pr.find(qn("w:pBdr")) is not None
        if paragraph.style.name == "Heading 1" and stripped not in STANDARD_DOCX_HEADINGS:
            errors.append(f"non-standard DOCX section heading found: {stripped!r}")
        if has_paragraph_border:
            if paragraph.style.name != "Heading 1" or stripped not in STANDARD_DOCX_HEADINGS:
                errors.append(f"paragraph border found outside standard section heading: {stripped!r}")
            else:
                bordered_headings.add(stripped)
        if paragraph.style.name == "List Bullet":
            errors.append(f"Word List Bullet paragraph found; use literal bullet glyph instead: {stripped!r}")
        if stripped.startswith("•"):
            literal_bullets += 1
            if not paragraph.text.startswith("•\t"):
                errors.append(f"literal bullet must use bullet + tab prefix: {stripped!r}")

    if literal_bullets == 0:
        errors.append("DOCX contains no literal bullet glyph paragraphs")

    missing_borders = STANDARD_DOCX_HEADINGS - bordered_headings
    if missing_borders:
        errors.append(f"standard section headings missing paragraph bottom border: {sorted(missing_borders)!r}")

    return errors


def check_docx_text(text: str, expectations: dict, resume: dict) -> list[str]:
    errors: list[str] = []
    errors.extend(check_required_exact(text, as_string_list(expectations.get("docx_required_exact"), "docx_required_exact")))
    errors.extend(check_order(text, as_string_list(expectations.get("docx_required_order"), "docx_required_order")))
    errors.extend(check_order(text, expected_document_order(resume, markdown=False)))
    errors.extend(check_forbidden(text, as_string_list(expectations.get("docx_forbidden"), "docx_forbidden")))
    errors.extend(check_yaml_content_present(text, resume))
    errors.extend(check_role_boundaries(text, resume))
    errors.extend(check_role_noun_starts(text))
    errors.extend(check_education_degree_fields(text, resume))
    return errors


def main() -> int:
    parser = argparse.ArgumentParser(description="Run structural and content checks against ATS DOCX.")
    parser.add_argument("--resume", type=Path, default=Path("resume.yaml"))
    parser.add_argument("--docx", type=Path, default=None)
    parser.add_argument("--expectations", type=Path, default=Path("tests/expected/ats_text_expectations.yaml"))
    args = parser.parse_args()

    try:
        docx_path = args.docx or default_artifact_path(args.resume, Path("output"), None, ".docx")
        document = Document(docx_path)
        text = docx_text(document)
        resume = load_resume(args.resume)
        expectations = load_expectations(args.expectations)
        errors = check_docx_structure(docx_path, document)
        errors.extend(check_docx_role_headings(document, resume))
        errors.extend(check_docx_text(text, expectations, resume))
    except (OSError, ValueError, zipfile.BadZipFile, yaml.YAMLError) as exc:
        print(f"DOCX check error: {exc}", file=sys.stderr)
        return 1

    if errors:
        print(f"DOCX checks failed for {docx_path}:", file=sys.stderr)
        for error in errors:
            print(f"  - {error}", file=sys.stderr)
        return 1

    print(f"DOCX checks passed for {docx_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
