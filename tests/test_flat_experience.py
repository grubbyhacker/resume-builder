from __future__ import annotations

import tempfile
import unittest
from contextlib import redirect_stdout
from io import StringIO
from pathlib import Path

import yaml
from docx import Document

from check_docx import check_docx_role_headings
from check_structured_markdown import check_structured_role_headings
from render_docx import render_docx
from render_structured_markdown import render_structured, structured_role_heading


ROOT = Path(__file__).resolve().parents[1]


class FlatExperienceProjectionTest(unittest.TestCase):
    def setUp(self) -> None:
        self.resume_path = ROOT / "resume.yaml"
        self.resume = yaml.safe_load(self.resume_path.read_text(encoding="utf-8"))

    def test_structured_markdown_has_one_complete_heading_per_role(self) -> None:
        text = render_structured(self.resume)

        self.assertEqual(check_structured_role_headings(text, self.resume), [])
        for job in self.resume["experience"]:
            self.assertNotIn(f"### {job['company']} - {job['location']}", text.splitlines())
            for role in job["roles"]:
                self.assertIn(f"### {structured_role_heading(job, role)}", text.splitlines())

    def test_structured_markdown_rejects_incomplete_role_heading(self) -> None:
        text = render_structured(self.resume)
        job = self.resume["experience"][1]
        role = job["roles"][1]
        complete = f"### {structured_role_heading(job, role)}"

        errors = check_structured_role_headings(text.replace(complete, f"### {role['title']}", 1), self.resume)

        self.assertTrue(any("complete structured Markdown role heading missing" in error for error in errors))

    def test_docx_has_one_complete_heading_per_role(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "resume.docx"
            with redirect_stdout(StringIO()):
                render_docx(self.resume_path, output_path, "20260713", "abcdef0")
            document = Document(output_path)

        self.assertEqual(check_docx_role_headings(document, self.resume), [])
        paragraphs = [(paragraph.text, paragraph.style.name) for paragraph in document.paragraphs]
        for job in self.resume["experience"]:
            self.assertNotIn((f"{job['company']} - {job['location']}", "Heading 2"), paragraphs)

    def test_docx_rejects_shared_employer_heading(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "resume.docx"
            with redirect_stdout(StringIO()):
                render_docx(self.resume_path, output_path, "20260713", "abcdef0")
            document = Document(output_path)

        job = self.resume["experience"][1]
        document.add_paragraph(f"{job['company']} - {job['location']}", style="Heading 2")

        errors = check_docx_role_headings(document, self.resume)

        self.assertTrue(any("shared employer heading" in error for error in errors))

    def test_docx_keeps_the_final_two_role_bullets_together(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "resume.docx"
            with redirect_stdout(StringIO()):
                render_docx(self.resume_path, output_path, "20260713", "abcdef0")
            document = Document(output_path)

        for job in self.resume["experience"]:
            for role in job["roles"]:
                bullets = role.get("bullets", [])
                if len(bullets) < 2:
                    continue
                penultimate = next(
                    paragraph for paragraph in document.paragraphs if paragraph.text.endswith(bullets[-2])
                )
                self.assertTrue(penultimate.paragraph_format.keep_with_next)


if __name__ == "__main__":
    unittest.main()
